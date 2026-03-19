import calendar
import datetime
from collections import defaultdict
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

class ExcelExporter:
    def generate(self, year, month, employees, all_shifts,
                 day_off_reqs=None, paid_leave_reqs=None, work_req_map=None,
                 hourly_groups=None, staffing_requirements=None, special_days=None):
        """
        シフトデータからExcelワークブックを生成する
        """
        # --- Guard against None ---
        if day_off_reqs is None: day_off_reqs = {}
        if paid_leave_reqs is None: paid_leave_reqs = {}
        if work_req_map is None: work_req_map = {}
        if hourly_groups is None: hourly_groups = {}
        if staffing_requirements is None: staffing_requirements = {}
        if special_days is None: special_days = {}

        wb = Workbook()
        ws = wb.active
        ws.title = f"{year}年{month}月"

        # --- Style Definitions ---
        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="4F81BD", end_color="4F81BD", fill_type="solid")
        center_alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        left_alignment = Alignment(horizontal='left', vertical='center')
        thin_border = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))
        thick_right_border = Border(left=Side(style='thin'), right=Side(style='thick'), top=Side(style='thin'), bottom=Side(style='thin'))

        red_font = Font(color="FF0000", bold=True)
        bold_font = Font(bold=True)
        saturday_fill = PatternFill(start_color="E0F0FF", end_color="E0F0FF", fill_type="solid")
        sunday_fill = PatternFill(start_color="FFECEC", end_color="FFECEC", fill_type="solid")
        deficit_fill = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid") # Yellow for deficit
        summary_header_fill = PatternFill(start_color="D3D3D3", end_color="D3D3D3", fill_type="solid") # Light grey
        prev_month_fill = PatternFill(start_color="F0F0F0", end_color="F0F0F0", fill_type="solid") # Very light grey for prev month context

        shift_colors = {
            "早": "C6EFCE", "日": "FFEB9C", "遅": "D9E1F2", "夜": "A5A5A5",
            "明": "FFC7CE", "休": "FFC7CE", "有": "FFC7CE"
        }
        
        # --- Data Pre-processing ---
        dates = sorted(list(set(s.date for s in all_shifts)))
        month_dates = [d for d in dates if d.year == year and d.month == month]
        shifts_by_user_date = {(s.user_id, s.date): s.shift_type.name for s in all_shifts}
        first_day_of_month = datetime.date(year, month, 1)

        # --- Main Header ---
        ws.cell(row=1, column=1, value="氏名").font = header_font
        ws.cell(row=1, column=1).fill = header_fill
        ws.cell(row=1, column=1).border = thin_border
        ws.column_dimensions['A'].width = 18

        for i, day in enumerate(dates):
            col = i + 2
            cell = ws.cell(row=1, column=col, value=f"{day.day}\n({day.strftime('%a')})")
            cell.font = header_font
            cell.alignment = center_alignment

            is_saturday = day.weekday() == 5
            is_sunday = day.weekday() == 6
            is_prev_month = day < first_day_of_month

            if is_saturday:
                cell.fill = PatternFill(start_color="6495ED", end_color="6495ED", fill_type="solid")
            elif is_sunday:
                cell.fill = PatternFill(start_color="DC143C", end_color="DC143C", fill_type="solid")
            else:
                cell.fill = header_fill
            
            if is_prev_month:
                cell.font = Font(bold=True, color="808080") # Grey out text
            
            # Add separator between previous and current month
            if i > 0 and dates[i-1].month != day.month:
                 for r in range(1, len(employees) + 2): # Apply to header and employee rows
                    ws.cell(row=r, column=col-1).border = thick_right_border
            else:
                cell.border = thin_border
            
            ws.column_dimensions[get_column_letter(col)].width = 5

        # --- Employee Summary Headers ---
        summary_start_col = len(dates) + 2
        summary_headers = ["勤務日", "休日", "その他休日", "夜勤"]
        for i, header in enumerate(summary_headers):
            col = summary_start_col + i
            cell = ws.cell(row=1, column=col, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.border = thin_border
            ws.column_dimensions[get_column_letter(col)].width = 6

        # --- Employee Rows ---
        for row_num, emp in enumerate(employees, start=2):
            ws.cell(row=row_num, column=1, value=emp.full_name).border = thin_border
            
            work_days, paid_holidays, holidays, night_shifts = 0, 0, 0, 0

            for i, day in enumerate(dates):
                col = i + 2
                shift_name = shifts_by_user_date.get((emp.id, day))
                
                is_month_date = day in month_dates
                if is_month_date:
                    if shift_name is None and day in paid_leave_reqs.get(emp.id, []):
                        shift_name = "有"
                    if shift_name:
                        if shift_name not in ["有", "休", "明"]: work_days += 1
                        if shift_name == "有": paid_holidays += 1
                        if shift_name in ["休", "明"]: holidays += 1
                        if "夜" in shift_name: night_shifts += 1

                cell = ws.cell(row=row_num, column=col, value=shift_name)
                cell.alignment = center_alignment
                
                # Add separator border
                if i > 0 and dates[i-1].month != day.month:
                    cell.border = Border(left=Side(style='thick'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))
                else:
                    cell.border = thin_border

                if (day in day_off_reqs.get(emp.id, []) and shift_name == "休") or \
                   (day in paid_leave_reqs.get(emp.id, []) and shift_name == "有") or \
                   ((emp.id, day) in work_req_map and shift_name in work_req_map.get((emp.id, day), [])):
                    cell.font = red_font
                
                if shift_name in shift_colors:
                    cell.fill = PatternFill(start_color=shift_colors[shift_name], end_color=shift_colors[shift_name], fill_type="solid")
                elif day.weekday() == 5: cell.fill = saturday_fill
                elif day.weekday() == 6: cell.fill = sunday_fill
                
                if day < first_day_of_month:
                    cell.fill = prev_month_fill


            emp_summary_data = [work_days, holidays, paid_holidays, night_shifts]
            for i, value in enumerate(emp_summary_data):
                ws.cell(row=row_num, column=summary_start_col + i, value=value).border = thin_border
                ws.cell(row=row_num, column=summary_start_col + i).alignment = center_alignment

        # --- Daily Summary Rows ---
        summary_start_row = len(employees) + 3
        
        # --- 時間帯別の人員配置計算 ---
        constraint_hours = [7, 8, 9, 12, 13, 14, 16, 18, 19]
        summary_labels = {h: f"{h:02d}:00時点" for h in constraint_hours}
        summary_labels["late_night_1"] = "20-23時"
        summary_labels["late_night_2"] = "23-24時"
        summary_labels["deep_night"] = "深夜勤(0-7時)"
        
        actual_counts = {label: {d: 0 for d in dates} for label in summary_labels.values()}
        
        shift_to_hours_map = defaultdict(dict)
        if hourly_groups:
            for hour, shifts_with_offset in hourly_groups.items():
                for s_name, offset in shifts_with_offset:
                    shift_to_hours_map[s_name][hour] = offset

        staff_by_category_date = {label: {d: set() for d in dates} for label in summary_labels.values()}
        
        full_shift_data = []
        for u in employees:
            for d in dates:
                s = shifts_by_user_date.get((u.id, d))
                if s: full_shift_data.append({'user_id': u.id, 'date': d, 'shift_name': s})

        for assignment in full_shift_data:
            shift_name = assignment["shift_name"]
            assignment_date = assignment["date"]
            emp_id = assignment["user_id"]

            if shift_name in ["休", "明", "有"]:
                continue
            
            covered_hours_info = shift_to_hours_map.get(shift_name)
            if not covered_hours_info:
                continue

            for hour, offset in covered_hours_info.items():
                # For deep night (0-6h), the count is for the day the shift *started*.
                if 0 <= hour <= 6:
                    target_date = assignment_date # No offset
                    if target_date in dates:
                        staff_by_category_date[summary_labels["deep_night"]][target_date].add(emp_id)

                # For all other times, the count is for the actual calendar day.
                else:
                    target_date = assignment_date + datetime.timedelta(days=offset)
                    if target_date not in dates:
                        continue
                    
                    if hour in constraint_hours:
                        staff_by_category_date[summary_labels[hour]][target_date].add(emp_id)
                    if 20 <= hour <= 22:
                        staff_by_category_date[summary_labels["late_night_1"]][target_date].add(emp_id)
                    if hour == 23:
                        staff_by_category_date[summary_labels["late_night_2"]][target_date].add(emp_id)

        for label, dates_data in staff_by_category_date.items():
            for d_date, emp_ids_set in dates_data.items():
                actual_counts[label][d_date] = len(emp_ids_set)

        # --- 集計行の書き込み ---
        ws.cell(row=summary_start_row, column=1, value="日付").border = thin_border
        ws.cell(row=summary_start_row, column=1).alignment = center_alignment
        
        for i, day in enumerate(dates):
            col = i + 2
            cell = ws.cell(row=summary_start_row, column=col, value=day.day)
            cell.border = thin_border
            cell.alignment = center_alignment
            if day < first_day_of_month:
                cell.fill = prev_month_fill
            if i > 0 and dates[i-1].month != day.month:
                 ws.cell(row=summary_start_row, column=col-1).border = thick_right_border
        
        ws.cell(row=summary_start_row + 1, column=1, value="項目").border = thin_border
        ws.cell(row=summary_start_row + 1, column=1).alignment = center_alignment
        
        for c in range(1, len(dates) + 2):
            ws.cell(row=summary_start_row, column=c).fill = summary_header_fill
            ws.cell(row=summary_start_row, column=c).font = bold_font
            ws.cell(row=summary_start_row + 1, column=c).fill = summary_header_fill
            ws.cell(row=summary_start_row + 1, column=c).font = bold_font


        current_row = summary_start_row + 2
        sorted_summary_keys = sorted(constraint_hours) + ["late_night_1", "late_night_2", "deep_night"]
        
        for key in sorted_summary_keys:
            is_time_key = isinstance(key, int)
            time_label = ""
            actual_row_values = []
            
            if is_time_key:
                hour_str_key = f"{key:02d}00"
                time_label = f"{key:02d}:00"
                actual_row_values = [actual_counts[summary_labels[key]][d] for d in dates]
            elif key == "late_night_1":
                hour_str_key = "2000_2300"; time_label = "20-23時"
                actual_row_values = [actual_counts[summary_labels["late_night_1"]][d] for d in dates]
            elif key == "late_night_2":
                hour_str_key = "2300_0000"; time_label = "23-24時"
                actual_row_values = [actual_counts[summary_labels["late_night_2"]][d] for d in dates]
            else: # deep_night
                hour_str_key = "0000_next_0700"; time_label = "0-7時"
                actual_row_values = [actual_counts[summary_labels["deep_night"]][d] for d in dates]
            
            req_weekday_key = f"min_staff_weekday_{hour_str_key}"
            req_sunday_key = f"min_staff_sunday_{hour_str_key}"
            
            base_req = {d: staffing_requirements.get(req_sunday_key, 0) if d.weekday() == 6 else staffing_requirements.get(req_weekday_key, 0) for d in dates}
            
            block_start_row = current_row
            
            def write_summary_row(row, label, values, is_deficit_check=False, is_final_req=False):
                ws.cell(row=row, column=1, value=label).border = thin_border
                for i, d in enumerate(dates):
                    col = i + 2
                    if d in month_dates:
                        val = values[i]
                        cell = ws.cell(row=row, column=col, value=f"+{val}" if (is_deficit_check and val > 0) else str(val))
                        cell.border = thin_border
                        if is_deficit_check and val < 0:
                            cell.fill = deficit_fill
                    else: # Previous month dates
                         ws.cell(row=row, column=col, value="").border = thin_border
                    # Separator
                    if i > 0 and dates[i-1].month != d.month:
                        ws.cell(row=row, column=col-1).border = thick_right_border
            
            if is_time_key and key in [8, 9]:
                num_visitors_list = [sum(sd.staff_increase for sd in special_days.get(d, []) if sd.visit_time and int(sd.visit_time.split(':')[0]) == key) for d in dates]
                final_req_list = [base_req[d] + num_visitors_list[i] for i, d in enumerate(dates)]
                diff_values = [actual - final for actual, final in zip(actual_row_values, final_req_list)]
                
                write_summary_row(current_row, "基本必要人数", [base_req[d] for d in dates])
                current_row += 1
                write_summary_row(current_row, "通院者数", num_visitors_list)
                current_row += 1
                write_summary_row(current_row, "最終必要人数", final_req_list, is_final_req=True)
                current_row += 1
            else:
                req_list = [base_req[d] for d in dates]
                diff_values = [actual - req for actual, req in zip(actual_row_values, req_list)]
                write_summary_row(current_row, "必要人数", req_list)
                current_row += 1

            write_summary_row(current_row, "予定人数", actual_row_values)
            current_row += 1
            write_summary_row(current_row, "過不足", diff_values, is_deficit_check=True)
            current_row += 1

            ws.merge_cells(start_row=block_start_row, start_column=1, end_row=current_row - 1, end_column=1)
            cell = ws.cell(row=block_start_row, column=1)
            cell.value = time_label
            cell.alignment = center_alignment
            cell.border = thin_border
        
        # Remove gridlines for a cleaner look
        ws.sheet_view.showGridLines = False

        return wb

